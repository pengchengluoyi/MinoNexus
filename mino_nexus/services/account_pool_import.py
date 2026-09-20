"""号池批量导入：粘贴/CSV 解析、预览、commit（upsert，不 replace_all）。"""
from __future__ import annotations

import csv
import io
import re
from typing import Any, Literal

from mino_nexus.services.account_ident_parse import _norm_phone

OnDuplicate = Literal["skip", "merge", "overwrite_credentials"]

_MAX_ROWS = 500
_PREVIEW_CAP = 80

# 表头别名 → 账号列或 facets.key
_HEADER_ALIASES: dict[str, str] = {
    "display_name": "display_name",
    "展示名": "display_name",
    "名称": "display_name",
    "备注名": "display_name",
    "phone": "phone",
    "手机": "phone",
    "手机号": "phone",
    "电话": "phone",
    "email": "email",
    "邮箱": "email",
    "username": "username",
    "用户名": "username",
    "账号": "username",
    "password": "password",
    "密码": "password",
    "otp": "otp",
    "验证码": "otp",
    "短信码": "otp",
    "note": "note",
    "备注": "note",
    "说明": "note",
    "env": "env",
    "环境": "env",
    "locked": "locked",
    "锁定": "locked",
    "手动占用": "locked",
    "account_id": "account_id",
    "id": "account_id",
    "lifecycle": "facet:lifecycle",
    "注册": "facet:lifecycle",
    "session": "facet:session",
    "登录态": "facet:session",
    "登录": "facet:session",
    "health": "facet:health",
    "健康": "facet:health",
    "profile_data": "facet:profile_data",
    "资料": "facet:profile_data",
}


def _slug_header(text: str) -> str:
    return re.sub(r"\s+", "", str(text or "").strip().lower())


def build_header_map(
    headers: list[str],
    pool_field_defs: list[dict[str, Any]],
) -> dict[int, str]:
    """列下标 → 目标字段（display_name / facet:key / …）。"""
    label_to_key: dict[str, str] = {}
    for d in pool_field_defs or []:
        key = str(d.get("key") or "").strip()
        if not key:
            continue
        label_to_key[_slug_header(key)] = f"facet:{key}"
        label_to_key[_slug_header(str(d.get("label") or ""))] = f"facet:{key}"

    out: dict[int, str] = {}
    for i, raw in enumerate(headers):
        h = str(raw or "").strip()
        if not h:
            continue
        key = _HEADER_ALIASES.get(h) or _HEADER_ALIASES.get(_slug_header(h))
        if not key:
            key = label_to_key.get(_slug_header(h))
        if key:
            out[i] = key
    return out


def _detect_delimiter(sample: str) -> str:
    first = ""
    for line in sample.splitlines():
        if line.strip():
            first = line
            break
    if "\t" in first:
        return "\t"
    if first.count(",") >= first.count(";"):
        return ","
    return ";"


def _cell_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, bool):
        return "true" if val else "false"
    if isinstance(val, float) and val == int(val):
        return str(int(val))
    return str(val).strip()


def parse_spreadsheet_bytes(filename: str, data: bytes) -> tuple[list[str], list[list[str]]]:
    """解析 .xlsx / .csv / .txt 表格，首行为表头。"""
    raw = bytes(data or b"")
    if not raw:
        return [], []
    name = str(filename or "").strip().lower()
    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("服务端未安装 openpyxl，无法解析 xlsx") from exc
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            ws = wb.active
            grid: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                cells = [_cell_str(c) for c in (row or ())]
                if any(cells):
                    grid.append(cells)
        finally:
            wb.close()
        if not grid:
            return [], []
        return grid[0], grid[1:]

    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return parse_table_text(raw.decode(enc))
        except UnicodeDecodeError:
            continue
    return [], []


def parse_table_text(text: str) -> tuple[list[str], list[list[str]]]:
    raw = str(text or "").strip()
    if not raw:
        return [], []
    delim = _detect_delimiter(raw)
    reader = csv.reader(io.StringIO(raw), delimiter=delim)
    rows = [[str(c or "").strip() for c in row] for row in reader]
    rows = [r for r in rows if any(cell for cell in r)]
    if not rows:
        return [], []
    headers = rows[0]
    body = rows[1:]
    return headers, body


def _parse_bool(val: str) -> bool | None:
    s = str(val or "").strip().lower()
    if not s:
        return None
    if s in ("1", "true", "yes", "是", "y", "锁定", "locked"):
        return True
    if s in ("0", "false", "no", "否", "n"):
        return False
    return None


def row_dict_from_cells(
    cells: list[str],
    col_map: dict[int, str],
    *,
    default_env: str = "test",
) -> dict[str, Any]:
    out: dict[str, Any] = {"env": default_env, "facets": {}}
    for idx, target in col_map.items():
        if idx >= len(cells):
            continue
        val = str(cells[idx] or "").strip()
        if not val:
            continue
        if target.startswith("facet:"):
            fk = target.split(":", 1)[1]
            out["facets"][fk] = val
            continue
        if target == "locked":
            b = _parse_bool(val)
            if b is not None:
                out["locked"] = b
            continue
        out[target] = val
    return out


def _match_key(row: dict[str, Any]) -> str:
    aid = str(row.get("account_id") or row.get("id") or "").strip()
    if aid:
        return f"id:{aid}"
    phone = _norm_phone(str(row.get("phone") or ""))
    if len(phone) >= 11:
        return f"phone:{phone}"
    email = str(row.get("email") or "").strip().lower()
    if email and "@" in email:
        return f"email:{email}"
    user = str(row.get("username") or "").strip().lower()
    if len(user) >= 2:
        return f"user:{user}"
    return ""


def _account_match_key(acc: dict[str, Any]) -> str:
    return _match_key(
        {
            "account_id": acc.get("id") or acc.get("account_id"),
            "phone": acc.get("phone"),
            "email": acc.get("email"),
            "username": acc.get("username"),
        }
    )


def _is_leased(acc: dict[str, Any]) -> bool:
    lease = acc.get("lease") if isinstance(acc.get("lease"), dict) else {}
    return bool(str(lease.get("run_id") or "").strip())


def _validate_row(row: dict[str, Any]) -> str:
    if not (
        str(row.get("display_name") or "").strip()
        or str(row.get("phone") or "").strip()
        or str(row.get("email") or "").strip()
        or str(row.get("username") or "").strip()
    ):
        return "至少填写展示名、手机、邮箱或用户名之一"
    return ""


def _apply_duplicate_policy(
    incoming: dict[str, Any],
    existing: dict[str, Any],
    *,
    on_duplicate: OnDuplicate,
) -> tuple[dict[str, Any] | None, str]:
    """返回 (payload, error)；None payload 表示 skip。"""
    if on_duplicate == "skip":
        return None, ""
    if _is_leased(existing):
        return None, "账号跑批占用中，已跳过"
    if existing.get("locked") and on_duplicate == "overwrite_credentials":
        return None, "账号已手动锁定，已跳过"

    old = dict(existing)
    aid = str(old.get("id") or old.get("account_id") or "")
    payload: dict[str, Any] = {"id": aid, "account_id": aid}

    if on_duplicate == "merge":
        for field in ("display_name", "phone", "email", "username", "note", "env"):
            new_v = str(incoming.get(field) or "").strip()
            old_v = str(old.get(field) or "").strip()
            payload[field] = new_v or old_v
        if str(incoming.get("password") or "").strip():
            payload["password"] = incoming["password"]
        elif str(old.get("password") or "").strip():
            payload["password"] = old["password"]
        if str(incoming.get("otp") or "").strip():
            payload["otp"] = incoming["otp"]
        elif str(old.get("otp") or "").strip():
            payload["otp"] = old["otp"]
        if incoming.get("locked") is not None:
            payload["locked"] = incoming["locked"]
        facets = dict(old.get("facets") or {})
        for k, v in (incoming.get("facets") or {}).items():
            if str(v or "").strip():
                facets[k] = v
        payload["facets"] = facets
        return payload, ""

    # overwrite_credentials
    for field in (
        "display_name",
        "phone",
        "email",
        "username",
        "password",
        "otp",
        "note",
        "env",
        "locked",
    ):
        if field in incoming and incoming.get(field) not in (None, ""):
            payload[field] = incoming[field]
        elif field in old:
            payload[field] = old[field]
    facets = dict(old.get("facets") or {})
    facets.update({k: v for k, v in (incoming.get("facets") or {}).items() if str(v or "").strip()})
    payload["facets"] = facets
    return payload, ""


def _load_parsed_account_rows(
    *,
    env_doc: dict[str, Any],
    text: str = "",
    rows: list[dict[str, Any]] | None = None,
    filename: str = "",
    file_data: bytes | None = None,
    default_env: str = "test",
) -> tuple[list[dict[str, Any]] | None, list[str], dict[str, str], str]:
    from mino_nexus.services.account_pool_templates import merged_pool_field_defs

    doc = env_doc if isinstance(env_doc, dict) else {}
    defs = merged_pool_field_defs(doc)
    if rows:
        return [dict(r) for r in rows if isinstance(r, dict)], [], {}, ""

    headers: list[str] = []
    body: list[list[str]] = []
    try:
        if file_data:
            headers, body = parse_spreadsheet_bytes(filename, file_data)
        elif str(text or "").strip():
            headers, body = parse_table_text(text)
        else:
            return None, [], {}, "请粘贴表格或上传文件"
    except ValueError as exc:
        return None, [], {}, str(exc)

    if not headers:
        return None, [], {}, "未解析到表头"
    col_map = build_header_map(headers, defs)
    if not col_map:
        return None, headers, {}, "无法识别列名，请使用手机号/邮箱等标准表头"
    col_map_out = {
        headers[i]: tgt for i, tgt in col_map.items() if i < len(headers)
    }
    parsed = [
        row_dict_from_cells(cells, col_map, default_env=default_env or "test")
        for cells in body[:_MAX_ROWS]
    ]
    return parsed, headers, col_map_out, ""


def preview_import(
    *,
    env_doc: dict[str, Any],
    project_id: str,
    text: str = "",
    rows: list[dict[str, Any]] | None = None,
    filename: str = "",
    file_data: bytes | None = None,
    default_env: str = "test",
    on_duplicate: OnDuplicate = "merge",
) -> dict[str, Any]:
    from mino_nexus.services.project_env import list_test_accounts

    pid = str(project_id or "").strip()
    doc = env_doc if isinstance(env_doc, dict) else {}
    existing = list_test_accounts(doc, project_id=pid)
    by_key: dict[str, dict[str, Any]] = {}
    for acc in existing:
        mk = _account_match_key(acc)
        if mk:
            by_key[mk] = acc

    parsed_rows, headers_out, col_map_out, load_err = _load_parsed_account_rows(
        env_doc=doc,
        text=text,
        rows=rows,
        filename=filename,
        file_data=file_data,
        default_env=default_env,
    )
    if load_err:
        return {
            "ok": False,
            "error": load_err,
            "headers": headers_out,
            "column_map": col_map_out,
            "preview": [],
            "stats": {},
            "source_filename": str(filename or ""),
        }
    parsed_rows = parsed_rows or []

    if len(parsed_rows) > _MAX_ROWS:
        return {
            "ok": False,
            "error": f"超过单次上限 {_MAX_ROWS} 行",
            "preview": [],
            "stats": {},
        }

    file_keys: set[str] = set()
    preview: list[dict[str, Any]] = []
    stats = {"create": 0, "update": 0, "skip": 0, "error": 0}

    for i, raw in enumerate(parsed_rows):
        row_num = i + 1
        err = _validate_row(raw)
        mk = _match_key(raw)
        if err:
            preview.append({"row": row_num, "action": "error", "error": err, "incoming": raw})
            stats["error"] += 1
            continue
        if mk and mk in file_keys:
            preview.append(
                {
                    "row": row_num,
                    "action": "error",
                    "error": "本文件内重复标识",
                    "incoming": raw,
                }
            )
            stats["error"] += 1
            continue
        if mk:
            file_keys.add(mk)

        hit = by_key.get(mk) if mk else None
        if hit:
            payload, pol_err = _apply_duplicate_policy(raw, hit, on_duplicate=on_duplicate)
            if pol_err:
                preview.append(
                    {
                        "row": row_num,
                        "action": "skip",
                        "reason": pol_err,
                        "account_id": hit.get("id"),
                        "incoming": raw,
                    }
                )
                stats["skip"] += 1
            elif payload is None:
                preview.append(
                    {
                        "row": row_num,
                        "action": "skip",
                        "reason": "已存在（跳过）",
                        "account_id": hit.get("id"),
                        "incoming": raw,
                    }
                )
                stats["skip"] += 1
            else:
                preview.append(
                    {
                        "row": row_num,
                        "action": "update",
                        "account_id": hit.get("id"),
                        "incoming": raw,
                    }
                )
                stats["update"] += 1
        else:
            preview.append({"row": row_num, "action": "create", "incoming": raw})
            stats["create"] += 1

    return {
        "ok": True,
        "headers": headers_out,
        "column_map": col_map_out,
        "preview": preview[:_PREVIEW_CAP],
        "preview_truncated": len(preview) > _PREVIEW_CAP,
        "stats": stats,
        "row_count": len(parsed_rows),
        "source_filename": str(filename or ""),
    }


def commit_import(
    *,
    env_doc: dict[str, Any],
    project_id: str,
    text: str = "",
    rows: list[dict[str, Any]] | None = None,
    filename: str = "",
    file_data: bytes | None = None,
    default_env: str = "test",
    on_duplicate: OnDuplicate = "merge",
) -> dict[str, Any]:
    from mino_nexus.services.account_pool_templates import new_account_id
    from mino_nexus.services.pool_account_store import persist_env_strip_test_accounts_if_needed
    from mino_nexus.services.project_env import list_test_accounts, save_one_test_account

    prev = preview_import(
        env_doc=env_doc,
        project_id=project_id,
        text=text,
        rows=rows,
        filename=filename,
        file_data=file_data,
        default_env=default_env,
        on_duplicate=on_duplicate,
    )
    if not prev.get("ok"):
        return {**prev, "committed": False}

    pid = str(project_id or "").strip()
    doc = dict(env_doc or {})
    existing = {str(a.get("id")): a for a in list_test_accounts(doc, project_id=pid)}
    by_key: dict[str, dict[str, Any]] = {}
    for acc in existing.values():
        mk = _account_match_key(acc)
        if mk:
            by_key[mk] = acc

    created: list[str] = []
    updated: list[str] = []
    skipped: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []

    source_rows, _, _, load_err = _load_parsed_account_rows(
        env_doc=doc,
        text=text,
        rows=rows,
        filename=filename,
        file_data=file_data,
        default_env=default_env,
    )
    if load_err:
        return {**prev, "ok": False, "error": load_err, "committed": False}
    source_rows = source_rows or []

    for i, raw in enumerate(source_rows):
        row_num = i + 1
        err = _validate_row(raw)
        if err:
            errors.append({"row": row_num, "error": err})
            continue
        mk = _match_key(raw)
        hit = by_key.get(mk) if mk else None
        try:
            if hit:
                payload, pol_err = _apply_duplicate_policy(raw, hit, on_duplicate=on_duplicate)
                if pol_err or payload is None:
                    skipped.append({"row": row_num, "reason": pol_err or "skip"})
                    continue
                saved = save_one_test_account(
                    doc,
                    payload,
                    project_id=pid,
                    account_id=str(hit.get("id") or ""),
                )
                aid = str(saved.get("id") or "")
                updated.append(aid)
                by_key[mk] = saved
                existing[aid] = saved
            else:
                aid = new_account_id()
                payload = dict(raw)
                payload["id"] = aid
                payload["account_id"] = aid
                saved = save_one_test_account(doc, payload, project_id=pid, account_id=aid)
                aid = str(saved.get("id") or aid)
                created.append(aid)
                mk2 = _account_match_key(saved)
                if mk2:
                    by_key[mk2] = saved
                existing[aid] = saved
        except ValueError as exc:
            errors.append({"row": row_num, "error": str(exc)})

    persist_env_strip_test_accounts_if_needed(pid, doc)
    return {
        "ok": True,
        "committed": True,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
        "stats": {
            "created": len(created),
            "updated": len(updated),
            "skipped": len(skipped),
            "errors": len(errors),
        },
    }
